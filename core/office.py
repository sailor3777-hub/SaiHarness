"""Creating Word and Excel files for the docx_write / xlsx_write tools (pure Python, no Qt).

Word is built from Markdown (markdown-it-py + python-docx): headings, paragraphs, bold/italic/code/links,
bulleted and numbered lists (nested), tables, quotes, code blocks, horizontal rules.
Excel is built from a "sheet name → rows" dict; a string value starting with "=" is written as a formula.
Both functions write atomically (temporary file + replace) and return (message, '') or (None, reason).
"""
import os
import shutil
from pathlib import Path

MAX_MARKDOWN_CHARS = 500_000
MAX_CELLS = 200_000
MAX_SHEETS = 30
BAD_SHEET_CHARS = set('[]:*?/\\')
CODE_FONT = "Consolas"


def _prepare_target(path: Path, overwrite: bool, suffixes: tuple[str, ...]) -> str:
    """An empty string if writing is allowed, otherwise the reason for refusal."""
    if path.suffix.lower() not in suffixes:
        return f"the file extension must be {' or '.join(suffixes)}"
    if path.is_dir():
        return "this path is a folder"
    if path.exists() and not overwrite:
        return "the file already exists; to replace it, pass overwrite=true (the old one will be kept as .bak)"
    return ""


def _atomic_save(path: Path, save, overwrite: bool) -> str:
    """Saves through save(tmp_path); an existing file is copied to .bak before being replaced. Returns the backup path or ''."""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".sai-tmp")
    backup = ""
    try:
        save(tmp)
        if path.exists() and overwrite:
            n, bak = 0, path.with_name(path.name + ".bak")
            while bak.exists():
                n += 1
                bak = path.with_name(f"{path.name}.bak.{n}")
            shutil.copy2(path, bak)
            backup = str(bak)
        os.replace(tmp, path)
    finally:
        try:
            tmp.unlink()
        except OSError:
            pass
    return backup


# ---------------------------------------------------------------- Word

def _add_inline(paragraph, children, base_bold=False, base_italic=False):
    bold, italic, link = base_bold, base_italic, None
    for tok in children:
        t = tok.type
        if t == "strong_open": bold = True
        elif t == "strong_close": bold = base_bold
        elif t == "em_open": italic = True
        elif t == "em_close": italic = base_italic
        elif t == "link_open": link = tok.attrGet("href")
        elif t == "link_close":
            if link:
                paragraph.add_run(f" ({link})")
            link = None
        elif t in ("softbreak", "hardbreak"):
            paragraph.add_run().add_break()
        elif t in ("text", "code_inline", "html_inline", "image"):
            text = tok.content
            if t == "image":
                text = tok.content or tok.attrGet("src") or ""
            run = paragraph.add_run(text)
            run.bold, run.italic = bold or None, italic or None
            if t == "code_inline":
                run.font.name = CODE_FONT


def _list_style(kind: str, depth: int) -> str:
    base = "List Bullet" if kind == "bullet" else "List Number"
    return base if depth <= 1 else f"{base} {min(depth, 3)}"


def markdown_to_docx(markdown: str):
    from docx import Document
    from docx.shared import Pt
    from markdown_it import MarkdownIt

    doc = Document()
    tokens = MarkdownIt("commonmark").enable("table").parse(markdown)
    lists: list[str] = []          # stack of open list kinds
    quote = 0                      # quote depth
    table: dict | None = None
    i = 0
    while i < len(tokens):
        tok = tokens[i]
        t = tok.type
        if t == "heading_open":
            inline = tokens[i + 1]
            p = doc.add_heading(level=min(int(tok.tag[1:]), 4))
            _add_inline(p, inline.children or [])
            i += 3
            continue
        if t == "paragraph_open":
            inline = tokens[i + 1]
            if table is not None:
                table["cell"].append(inline.content)
            else:
                style = _list_style(lists[-1], len(lists)) if lists else ("Quote" if quote else None)
                p = doc.add_paragraph(style=style)
                _add_inline(p, inline.children or [])
            i += 3
            continue
        if t in ("bullet_list_open", "ordered_list_open"):
            lists.append("bullet" if t == "bullet_list_open" else "ordered")
        elif t in ("bullet_list_close", "ordered_list_close"):
            lists.pop()
        elif t == "blockquote_open":
            quote += 1
        elif t == "blockquote_close":
            quote -= 1
        elif t in ("fence", "code_block"):
            p = doc.add_paragraph()
            run = p.add_run(tok.content.rstrip("\n"))
            run.font.name, run.font.size = CODE_FONT, Pt(9)
        elif t == "hr":
            doc.add_paragraph("―" * 20)
        elif t == "table_open":
            table = {"rows": [], "row": None, "cell": None, "head": 0}
        elif t == "tr_open":
            table["row"] = []
        elif t in ("th_open", "td_open"):
            table["cell"] = []
            if t == "th_open":
                table["head"] = len(table["rows"]) + 1
        elif t in ("th_close", "td_close"):
            table["row"].append(" ".join(table["cell"]))
            table["cell"] = None
        elif t == "inline" and table is not None and table["cell"] is not None:
            table["cell"].append(tok.content)
        elif t == "tr_close":
            table["rows"].append(table["row"])
        elif t == "table_close":
            rows = table["rows"]
            width = max((len(r) for r in rows), default=0)
            if rows and width:
                grid = doc.add_table(rows=len(rows), cols=width)
                grid.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, text in enumerate(row):
                        cell = grid.cell(r, c)
                        run = cell.paragraphs[0].add_run(text)
                        if r < table["head"]:
                            run.bold = True
                doc.add_paragraph()
            table = None
        i += 1
    return doc


def write_docx(markdown: str, path: Path, overwrite: bool = False) -> tuple[str | None, str]:
    if not markdown.strip():
        return None, "markdown is empty"
    if len(markdown) > MAX_MARKDOWN_CHARS:
        return None, f"the text is too long (over {MAX_MARKDOWN_CHARS} characters)"
    reason = _prepare_target(path, overwrite, (".docx",))
    if reason:
        return None, reason
    try:
        doc = markdown_to_docx(markdown)
        backup = _atomic_save(path, lambda tmp: doc.save(str(tmp)), overwrite)
    except OSError as e:
        return None, f"writing failed ({e.__class__.__name__}: {e.strerror or e})"
    extra = f" Copy of the previous file: {backup}." if backup else ""
    return f"Done: created the Word document {path}.{extra}", ""


# ---------------------------------------------------------------- Excel

def write_xlsx(sheets: dict, path: Path, overwrite: bool = False) -> tuple[str | None, str]:
    from openpyxl import Workbook
    if not isinstance(sheets, dict) or not sheets:
        return None, "sheets must be a non-empty object {sheet name: [[row cells], ...]}"
    if len(sheets) > MAX_SHEETS:
        return None, f"more than {MAX_SHEETS} sheets"
    cells = 0
    for name, rows in sheets.items():
        if not name or len(name) > 31 or BAD_SHEET_CHARS & set(name):
            return None, f"invalid sheet name \"{name}\" (up to 31 characters, without []:*?/\\)"
        if not isinstance(rows, list) or any(not isinstance(r, list) for r in rows):
            return None, f"sheet \"{name}\": a list of rows is expected, each row a list of cells"
        cells += sum(len(r) for r in rows)
    if cells > MAX_CELLS:
        return None, f"more than {MAX_CELLS} cells"
    reason = _prepare_target(path, overwrite, (".xlsx",))
    if reason:
        return None, reason
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append([c if isinstance(c, (str, int, float, bool)) or c is None else str(c) for c in row])
    try:
        backup = _atomic_save(path, lambda tmp: wb.save(str(tmp)), overwrite)
    except OSError as e:
        return None, f"writing failed ({e.__class__.__name__}: {e.strerror or e})"
    extra = f" Copy of the previous file: {backup}." if backup else ""
    return f"Done: created the Excel workbook {path} (sheets: {len(sheets)}, cells: {cells}).{extra}", ""
