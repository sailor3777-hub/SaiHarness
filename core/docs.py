"""Text from PDF, Word and Excel for the fs_read tool (pure Python, no Qt).

read_document(path) returns (text, '') or (None, reason). Formats: .pdf (pypdf), .docx (python-docx),
.xlsx/.xlsm (openpyxl). Old .doc/.xls are not supported. The result is not cut to the model's size:
the caller does that (fs_read returns the text in chunks with offset).
"""
import os
import threading
from collections import OrderedDict

from core.markers import NOT_FOUND

DOC_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xlsm"}
MAX_DOC_BYTES = 30_000_000     # a larger file is not opened
MAX_DOC_CHARS = 2_000_000      # no text is extracted beyond this
MAX_SHEET_ROWS = 20_000


def is_document_name(name: str) -> bool:
    return os.path.splitext(name)[1].lower() in DOC_EXTENSIONS


class _Out:
    """A text accumulator with a size limit."""

    def __init__(self):
        self.parts: list[str] = []
        self.size = 0
        self.full = False

    def add(self, s: str) -> bool:
        if self.full:
            return False
        self.parts.append(s)
        self.size += len(s) + 1
        if self.size >= MAX_DOC_CHARS:
            self.full = True
        return not self.full

    def text(self) -> str:
        t = "\n".join(self.parts)
        if self.full:
            t = t[:MAX_DOC_CHARS] + f"\n[the document is longer than {MAX_DOC_CHARS} characters, the rest was not extracted]"
        return t


def _read_pdf(path: str) -> tuple[str | None, str]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError
    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            try:
                ok = reader.decrypt("")  # a PDF with an empty password opens without asking
            except Exception:
                ok = 0
            if not ok:
                return None, "the PDF is password protected"
        out = _Out()
        pages = len(reader.pages)
        has_text = False
        for i, page in enumerate(reader.pages, 1):
            text = (page.extract_text() or "").strip()
            has_text = has_text or bool(text)
            if not out.add(f"--- page {i} of {pages} ---\n{text}" if text else f"--- page {i} of {pages} --- (no text)"):
                break
    except PdfReadError as e:
        return None, f"the PDF is damaged ({e})"
    if not has_text:
        return None, "the PDF has no text layer (looks like a scan), it cannot be read as text"
    return out.text(), ""


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).replace("\t", " ").replace("\r", " ").replace("\n", " ")


def _read_docx(path: str) -> tuple[str | None, str]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    try:
        document = docx.Document(path)
    except Exception as e:
        return None, f"does not open as Word ({type(e).__name__})"
    out = _Out()
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "") if para.style is not None else ""
            # style names are checked in English and Russian: files saved by a Russian Word keep localized names
            if style.startswith(("Heading", "Заголовок")):
                digits = "".join(c for c in style if c.isdigit())
                text = "#" * (int(digits) if digits else 1) + " " + text
            elif style == "Title":
                text = "# " + text
            elif "List" in style or "Список" in style:
                text = "- " + text
            if not out.add(text):
                break
        elif tag == "tbl":
            table = Table(child, document)
            ok = True
            for row in table.rows:
                cells, last_tc = [], None
                for c in row.cells:
                    if c._tc is not last_tc:  # a merged cell repeats in row.cells
                        cells.append(_cell(c.text.strip()))
                        last_tc = c._tc
                ok = out.add(" | ".join(cells))
                if not ok:
                    break
            if not ok:
                break
            out.add("")
    return out.text(), ""


def _read_xlsx(path: str) -> tuple[str | None, str]:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        return None, f"does not open as Excel ({type(e).__name__})"
    out = _Out()
    try:
        for ws in wb.worksheets:
            rows = 0
            if not out.add(f"=== Sheet \"{ws.title}\" ==="):
                break
            for row in ws.iter_rows(values_only=True):
                cells = [_cell(v) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if not cells:
                    continue
                rows += 1
                if rows > MAX_SHEET_ROWS:
                    out.add(f"[more than {MAX_SHEET_ROWS} rows, the rest were not extracted]")
                    break
                if not out.add("\t".join(cells)):
                    break
            if out.full:
                break
    finally:
        wb.close()
    return out.text(), ""


_READERS = {".pdf": _read_pdf, ".docx": _read_docx, ".xlsx": _read_xlsx, ".xlsm": _read_xlsx}


def read_document(path: str) -> tuple[str | None, str]:
    """Document text or (None, reason). Excel formulas are given as the last saved value."""
    ext = os.path.splitext(path)[1].lower()
    reader = _READERS.get(ext)
    if reader is None:
        return None, "format not supported"
    try:
        size = os.path.getsize(path)
    except FileNotFoundError:
        return None, NOT_FOUND
    except OSError:
        return None, "no access to the file"
    if size > MAX_DOC_BYTES:
        return None, f"too large ({size // 1024 // 1024} MB)"
    try:
        return reader(path)
    except PermissionError:
        return None, "no access to the file"
    except Exception as e:  # broken files raise all kinds of library exceptions
        return None, f"could not be read ({type(e).__name__})"


_cache: "OrderedDict[tuple, tuple[str | None, str]]" = OrderedDict()
_cache_lock = threading.Lock()
_CACHE_SIZE = 8


def read_document_cached(path: str) -> tuple[str | None, str]:
    """read_document cached by (path, modification time, size): the Files panel reads the selection on every summary and send."""
    try:
        st = os.stat(path)
    except OSError:
        return read_document(path)
    key = (os.path.abspath(path), st.st_mtime_ns, st.st_size)
    with _cache_lock:
        if key in _cache:
            _cache.move_to_end(key)
            return _cache[key]
    result = read_document(path)
    with _cache_lock:
        _cache[key] = result
        while len(_cache) > _CACHE_SIZE:
            _cache.popitem(last=False)
    return result
