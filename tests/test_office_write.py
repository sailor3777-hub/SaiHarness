"""docx_write / xlsx_write on the real MCP server: content, overwriting, refusals; the result is read back with core.docs."""
import sys, pathlib, tempfile
from core import db, docs
from core.mcp_manager import load_config
from model_view import McpManager  # errors with the prefix, as the model sees them

ok_all = True
def check(name, cond, extra=""):
    global ok_all
    ok_all &= bool(cond)
    print(("OK   " if cond else "FAIL ") + name + (f"  [{extra}]" if extra and not cond else ""))

tmp = pathlib.Path(tempfile.mkdtemp()); db.DB_PATH = tmp / "office.db"
root = tmp / "proj"; root.mkdir(); db.add_root(str(root))
m = McpManager(); m.connect_all(load_config(tmp / "mcp.json", db.DB_PATH))
def call(tool, **a): return m.call("toolbox", tool, a, 60)[0]
tools = {t.name: t for t in m.all_tools()}
check("both tools write (they will ask)", not tools["docx_write"].read_only and not tools["xlsx_write"].read_only)

MD = """# Report

The first paragraph with **bold**, *italic* and `code`, a link [site](https://example.com).

## List

- item one
- item two
  - nested
1. first
2. second

| Name | Amount |
|-----|-------|
| Anna | 10 |
| Boris | 20 |

> a quote

```
print("code")
```
"""
d = root / "report.docx"
t = call("docx_write", path=str(d), markdown=MD)
check("docx created", t.startswith("Done") and d.exists() and d.stat().st_size > 2000, t)
text, reason = docs.read_document(str(d))
check("read back through docs", text is not None and reason == "", reason)
check("headings as headings", "# Report" in text and "## List" in text, text)
check("paragraph and link", "bold" in text and "italic" in text and "https://example.com" in text, text)
check("lists", "item one" in text and "nested" in text and "second" in text, text)
check("a table as rows", "Name | Amount" in text and "Anna | 10" in text and "Boris | 20" in text, text)
check("quote and code", "a quote" in text and 'print("code")' in text, text)
import docx as _docx
dd = _docx.Document(str(d))
check("bold is really bold", any(r.bold and r.text == "bold" for p in dd.paragraphs for r in p.runs))
check("a table as a 3x2 table", len(dd.tables) == 1 and len(dd.tables[0].rows) == 3 and len(dd.tables[0].columns) == 2)
check("no temporary files left", [x.name for x in root.iterdir() if x.name.endswith(".sai-tmp")] == [])

# ---- protection against overwriting ----
before = d.read_bytes()
t = call("docx_write", path=str(d), markdown="# New")
check("an existing file is not overwritten", t.startswith("ERROR") and "overwrite" in t and d.read_bytes() == before, t)
t = call("docx_write", path=str(d), markdown="# New", overwrite=True)
bak = root / "report.docx.bak"
check("overwrite=true: replaced, the old one in .bak", t.startswith("Done") and "Copy of the previous" in t and bak.exists() and bak.read_bytes() == before, t)
check("the new content", "New" in docs.read_document(str(d))[0])
t = call("docx_write", path=str(d), markdown="# Another", overwrite=True)
check("the second backup does not overwrite the first (.bak.1)", (root / "report.docx.bak.1").exists() and bak.read_bytes() == before, t)

# ---- refusals ----
check("not .docx", call("docx_write", path=str(root / "a.txt"), markdown="x").startswith("ERROR"))
check("empty markdown", "empty" in call("docx_write", path=str(root / "e.docx"), markdown="  "))
(root / "dir.docx").mkdir(); check("a folder instead of a file", "folder" in call("docx_write", path=str(root / "dir.docx"), markdown="x"))
rel = call("docx_write", path="rel.docx", markdown="# Relative")
check("a relative path: from the first allowed folder", rel.startswith("Done") and (root / "rel.docx").exists(), rel)
sub = call("docx_write", path=str(root / "new" / "deep" / "x.docx"), markdown="# In a new folder")
check("missing folders are created", sub.startswith("Done") and (root / "new" / "deep" / "x.docx").exists(), sub)

# ---- Excel ----
x = root / "tab.xlsx"
t = call("xlsx_write", path=str(x), sheets={"Data": [["Name", "Amount"], ["Anna", 10], ["Boris", 20.5], ["Total", "=SUM(B2:B3)"]],
                                            "Empty": [], "Flags": [[True, None, "text"]]})
check("xlsx created", t.startswith("Done") and "sheets: 3" in t and x.exists(), t)
text, reason = docs.read_document(str(x))
check("read back through docs", text is not None and "Anna" in text and "Boris" in text and "20.5" in text, reason or text)
import openpyxl
wb = openpyxl.load_workbook(str(x))
check("sheets and names", wb.sheetnames == ["Data", "Empty", "Flags"], wb.sheetnames)
check("a formula is written as a formula", wb["Data"]["B4"].value == "=SUM(B2:B3)" and wb["Data"]["B2"].value == 10)
before = x.read_bytes()
t = call("xlsx_write", path=str(x), sheets={"A": [[1]]})
check("xlsx: an existing file is not overwritten", t.startswith("ERROR") and x.read_bytes() == before, t)
t = call("xlsx_write", path=str(x), sheets={"A": [[1]]}, overwrite=True)
check("xlsx: overwrite makes a .bak", t.startswith("Done") and (root / "tab.xlsx.bak").read_bytes() == before, t)
check("a bad sheet name", "sheet name" in call("xlsx_write", path=str(root / "b.xlsx"), sheets={"a/b": [[1]]}))
check("a long sheet name", "sheet name" in call("xlsx_write", path=str(root / "b.xlsx"), sheets={"x" * 32: [[1]]}))
check("not .xlsx", call("xlsx_write", path=str(root / "b.csv"), sheets={"a": [[1]]}).startswith("ERROR"))
check("empty sheets", call("xlsx_write", path=str(root / "b.xlsx"), sheets={}).startswith("ERROR"))
check("no file is created on refusal", not (root / "b.xlsx").exists())

m.shutdown()
print("ALL OK" if ok_all else "SOME FAILED")
