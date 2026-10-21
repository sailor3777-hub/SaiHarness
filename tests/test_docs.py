"""Reading PDF / Word / Excel: the core.docs core and the fs_read tool through the real toolbox MCP server (including offset)."""
import pathlib, tempfile
import docx, openpyxl
from docx.enum.style import WD_STYLE_TYPE
from pypdf import PdfWriter
from core import db, docs
from core.mcp_manager import default_config
from model_view import McpManager  # errors with the prefix, as the model sees them

ok_all = True
def check(name, cond, extra=""):
    global ok_all
    ok_all &= bool(cond)
    print(("OK   " if cond else "FAIL ") + name + (f"  [{extra}]" if extra and not cond else ""))

tmp = pathlib.Path(tempfile.mkdtemp())
db.DB_PATH = tmp / "saiharness.db"
proj = tmp / "proj"; proj.mkdir()
db.add_root(str(proj))


def make_pdf(path, pages):
    """A minimal PDF with text (Helvetica, Latin only): one page per element of pages."""
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font_id = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
                    f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode())
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = bytearray(b"%PDF-1.4\n"); offsets = []
    for n, body in enumerate(objs, 1):
        offsets.append(len(out)); out += b"%d 0 obj\n" % n + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    path.write_bytes(bytes(out))


# ---------- PDF ----------
make_pdf(proj / "doc.pdf", ["Hello PDF page one", "Second page MAGIC-77"])
t, r = docs.read_document(str(proj / "doc.pdf"))
check("PDF: text and page numbers", t and r == "" and "--- page 1 of 2 ---" in t and "Hello PDF page one" in t and "MAGIC-77" in t, (t, r))
w = PdfWriter(); w.add_blank_page(200, 200)
with open(proj / "scan.pdf", "wb") as f: w.write(f)
t, r = docs.read_document(str(proj / "scan.pdf"))
check("PDF without text: a clear reason", t is None and "no text layer" in r, (t, r))
w = PdfWriter(); w.add_blank_page(200, 200); w.encrypt("secret")
with open(proj / "locked.pdf", "wb") as f: w.write(f)
t, r = docs.read_document(str(proj / "locked.pdf"))
check("PDF with a password: the reason", t is None and "password" in r, (t, r))
(proj / "broken.pdf").write_bytes(b"%PDF-1.4 not really")
t, r = docs.read_document(str(proj / "broken.pdf"))
check("a broken PDF: a reason, no crash", t is None and r, (t, r))

# ---------- Word ----------
d = docx.Document()
d.add_heading("Project report", 1)
d.add_paragraph("The first paragraph with the number 42.")
d.add_heading("Subsection", 2)
d.add_paragraph("a list item", style="List Bullet")
tb = d.add_table(rows=2, cols=2)
tb.cell(0, 0).text = "Name"; tb.cell(0, 1).text = "Amount"; tb.cell(1, 0).text = "Rent"; tb.cell(1, 1).text = "1500"
d.add_paragraph("Conclusion AMBER-9.")
d.add_paragraph("Heading from a Russian Word", style=d.styles.add_style("Заголовок 2", WD_STYLE_TYPE.PARAGRAPH))   # a localized style name
d.save(proj / "report.docx")
t, r = docs.read_document(str(proj / "report.docx"))
check("Word: headings", t and "# Project report" in t and "## Subsection" in t, t)
check("Word: a heading with a Russian style name (files from a Russian Word)", "## Heading from a Russian Word" in t, t)
check("Word: paragraphs and a list", "The first paragraph with the number 42." in t and "- a list item" in t, t)
check("Word: a table as rows", "Name | Amount" in t and "Rent | 1500" in t, t)
check("Word: the order as in the document", t.index("a list item") < t.index("Rent | 1500") < t.index("AMBER-9"), t)
(proj / "fake.docx").write_bytes(b"not a zip")
t, r = docs.read_document(str(proj / "fake.docx"))
check("a broken docx: the reason", t is None and "Word" in r, (t, r))

# ---------- Excel ----------
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Budget"
ws.append(["Item", "Amount"]); ws.append(["Rent", 1500.0]); ws.append([None, None]); ws.append(["Total", "=B2*2"])
ws2 = wb.create_sheet("Empty")
wb.create_sheet("Data").append([1, 2.5, "x\ty"])
wb.save(proj / "book.xlsx")
t, r = docs.read_document(str(proj / "book.xlsx"))
check("Excel: sheets and rows", t and "=== Sheet \"Budget\" ===" in t and "Item\tAmount" in t and "Rent\t1500" in t, t)
check("Excel: empty rows skipped, floats without \".0\"", "\n\n" not in t.split("=== Sheet \"Data\"")[0].strip() and "1500.0" not in t, t)
check("Excel: a formula without a saved value does not crash", "Total" in t, t)
check("Excel: a tab in a cell does not break the columns", "1\t2.5\tx y" in t, t)
check("Excel: the second sheet and an empty sheet", "=== Sheet \"Empty\" ===" in t and "=== Sheet \"Data\" ===" in t, t)
t, r = docs.read_document(str(proj / "nope.xlsx"))
check("no file", t is None and r == "file not found", (t, r))
t, r = docs.read_document(str(proj / "x.doc"))
check("an old format is not supported", t is None and "not supported" in r, (t, r))

# ---------- fs_read through MCP ----------
m = McpManager(log_dir=tmp / "logs"); m.connect_all(default_config(db.DB_PATH))
txt, _ = m.call("toolbox", "fs_read", {"path": "doc.pdf"}, 60)
check("fs_read PDF", "MAGIC-77" in txt and "ERROR" not in txt, txt)
txt, _ = m.call("toolbox", "fs_read", {"path": "report.docx"}, 60)
check("fs_read docx", "AMBER-9" in txt, txt)
txt, _ = m.call("toolbox", "fs_read", {"path": "book.xlsx"}, 60)
check("fs_read xlsx", "Rent\t1500" in txt, txt)
txt, _ = m.call("toolbox", "fs_read", {"path": "scan.pdf"}, 60)
check("fs_read of a scan: an error with the reason", "ERROR" in txt and "text layer" in txt, txt)
txt, _ = m.call("toolbox", "fs_read", {"path": "old.xls"}, 60)
check("fs_read of a missing .xls: file not found", "ERROR" in txt, txt)

# chunks through offset
big = "".join(f"line {i:05d}\n" for i in range(3000))
(proj / "big.txt").write_bytes(big.encode("utf-8"))
got, off, parts = "", 0, 0
while True:
    txt, _ = m.call("toolbox", "fs_read", {"path": "big.txt", "offset": off}, 60)
    parts += 1
    if "[truncated:" not in txt:
        body = txt.split("\n[shown from character")[0]
        got += body
        break
    body, tail = txt.rsplit("\n[truncated:", 1)
    got += body
    off = int(tail.split("offset=")[1].rstrip("]"))
    if parts > 20: break
check("offset: the chunks join into the original text", got == big and parts == len(big) // 10_500 + 1, (parts, len(got), len(big)))
check("a chunk fits into the router's limit", all(len(x) < 12_000 for x in [txt]), len(txt))
txt, _ = m.call("toolbox", "fs_read", {"path": "big.txt", "offset": 10_000_000}, 60)
check("offset past the end: an error", "ERROR" in txt and "past the end" in txt, txt)
txt, _ = m.call("toolbox", "fs_read", {"path": "big.txt", "max_chars": 50_000}, 60)
check("max_chars above the limit is reduced", len(txt) < 11_000 and "offset=10500" in txt, len(txt))
m.shutdown()

print("\nALL OK" if ok_all else "\nSOME FAILED")
raise SystemExit(0 if ok_all else 1)
