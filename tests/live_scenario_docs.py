"""A live run of reading PDF / Word / Excel through fs_read: the facts are in the documents, the model's answer is checked against them.
Usage: python tests/live_scenario_docs.py <model> <max_tokens>"""
import json, sys, time, pathlib, tempfile
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication, QMessageBox
from core import db
tmp = pathlib.Path(tempfile.mkdtemp()); db.DB_PATH = tmp / "t44.db"
from core.context_builder import estimate_request_tokens
from ui.main_window import MainWindow

MODEL = sys.argv[1]; MAXTOK = int(sys.argv[2])
app = QApplication([])
def wait(cond, timeout=900.0):
    t = time.time()
    while not cond() and time.time() - t < timeout: app.processEvents(); time.sleep(0.01)
    app.processEvents(); return cond()


import docx, openpyxl
from pypdf import PdfWriter
def make_pdf(path, pages):
    """A minimal PDF with text (Helvetica, Latin only): one page per element of pages."""
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font_id = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
                    f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode())
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = bytearray(b"%PDF-1.4\n"); offsets = []
    for n, body in enumerate(objs, 1):
        offsets.append(len(out)); out += b"%d 0 obj\n" % n + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    path.write_bytes(bytes(out))



d = tmp / "Docs"; d.mkdir()
make_pdf(d / "contract.pdf", ["Contract between ACME and Beta. Term: 12 months.", "Delivery code: KAPPA-731. Penalty: 5 percent."])
w0 = PdfWriter(); w0.add_blank_page(300, 300)
with open(d / "scan.pdf", "wb") as f: w0.write(f)
dd = docx.Document(); dd.add_heading("Repair report", 1); dd.add_paragraph("The office repair was finished in March.")
tb = dd.add_table(rows=3, cols=2); rows = [("Item", "Amount, USD"), ("Materials", "48250"), ("Labour", "31000")]
for i, (a, b) in enumerate(rows): tb.cell(i, 0).text = a; tb.cell(i, 1).text = b
dd.add_paragraph("Responsible person: Ivanova A. P."); dd.save(d / "report.docx")
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Budget"
for r in (["Month", "Expense"], ["January", 1200], ["February", 3400], ["March", 560]): ws.append(r)
wb.save(d / "budget.xlsx")
long = openpyxl  # placeholder to keep names tidy
ld = docx.Document(); ld.add_heading("Long regulation", 1)
for i in range(1, 400): ld.add_paragraph(f"Item {i}. Standard provisions of the regulation, nothing special here.")
ld.add_paragraph("The secret code word at the end of the regulation: LAPIS-58."); ld.save(d / "long.docx")
db.add_root(str(d))
asked = []
QMessageBox.question = staticmethod(lambda p, t, text, *a, **k: (asked.append(text), QMessageBox.Yes)[1])

w = MainWindow(mcp_config=tmp / "mcp.json"); tp = w.tools_panel; v = w.chat_view; ctx = w.context; fp = w.files_panel
wait(lambda: tp.states and not tp.connect_req.busy, 40)
c = ctx.connections; c.name.setText("lm"); c.url.setText("127.0.0.1:1234"); c.save_btn.click()
b = ctx.builder; wait(lambda: not b.models_req.busy)
b.name.setText("p"); b.base_model.setCurrentText(MODEL); b.max_tokens.setValue(MAXTOK); b.temperature.setValue(0.7); b.max_steps.setValue(16); b.save_btn.click()
tp.tree.topLevelItem(0).setCheckState(0, Qt.Checked)
wait(lambda: not w.limit_req.busy, 20)
print(f"\n##### {MODEL}; max_tokens {MAXTOK}; context limit reported by the server: {w._limit_info}; catalog ≈{tp.prompt_tokens()} tok.")

usage_log = []
orig_usage = w._on_usage
def logging_usage(sid, prompt, completion, raw):
    usage_log.append((prompt, completion, raw)); orig_usage(sid, prompt, completion, raw)
w._on_usage = logging_usage
# thread.usage is already connected to the original bound method when the thread is created: it is replaced through the wrapper below
orig_submit = w.on_submit
import worker
orig_init = worker.ChatStreamer.__init__
def patched_init(self, *a, **k):
    orig_init(self, *a, **k); self.usage.connect(lambda sid, p, c_, e: usage_log.append((p, c_, e)))
worker.ChatStreamer.__init__ = patched_init

def ask(question):
    usage_log.clear(); n0 = len(db.list_messages(w.chat_id)) if w.chat_id else 0
    t0 = time.time(); v.input.setPlainText(question); v.send_btn.click()
    done = wait(lambda: w.stream is None)
    rows = db.list_messages(w.chat_id)[n0:]
    calls = [x for r in rows if r["role"] == "assistant" and r["tool_calls"] for x in json.loads(r["tool_calls"])]
    names = [(json.loads(x["function"]["arguments"] or "{}").get("tool") or x["function"]["name"]) for x in calls]
    errors = [r for r in rows if r["role"] == "tool" and r["content"].startswith("ERROR")]
    final = next((r for r in reversed(rows) if r["role"] == "assistant" and not r["tool_calls"]), None)
    sysrows = [r["content"] for r in rows if r["role"] == "system"]
    print(f"\n=== {question[:110]!r}\n    {time.time()-t0:.0f} s, finished: {done}; calls: {len(calls)} ({', '.join(names)}); tool errors: {len(errors)}")
    if usage_log:
        print("    request tokens (server / our raw estimate): " + "; ".join(f"{p}/{e}" for p, _, e in usage_log) + f"; correction ×{w._scale.get(w._limit_key, 1.0):.2f}")
    for r in errors: print("    error:", r["content"][:140].replace("\n", " "))
    for s in sysrows: print("    system:", s)
    print("    answer:", (final["content"] if final else "(none)")[:380].replace("\n", " "))
    return ((final["content"] or final["reasoning"] or "") if final else "", names, errors)  # Qwen 3.5 sometimes puts the answer in reasoning


verdicts = []
def expect(label, final_text, *needles):
    ok = all(n.lower() in (final_text or "").lower() for n in needles)
    verdicts.append((label, ok)); print(f"    CHECK {label}: {'OK' if ok else 'FAIL'} (expected: {', '.join(needles)})")

ra, names, _ = ask("Read the file contract.pdf and tell me the delivery code and which page it is on.")
expect("PDF: delivery code", ra, "KAPPA-731"); expect("PDF: fs_read was called", " ".join(names), "fs_read")
ra, names, _ = ask("Open report.docx and give the amount for materials and the surname of the person responsible.")
expect("Word: table and paragraph", "".join(ch for ch in ra if not ch.isspace()), "48250", "Ivanova")
ra, names, _ = ask("budget.xlsx has expenses by month. Add up the total expense with the calculator and give the number.")
expect("Excel: total expense 5160", "".join(ch for ch in ra if not ch.isspace()), "5160")
ra, names, _ = ask("The file long.docx is long. Read it to the end and give the secret code word from the last paragraph.")
expect("long Word: reading in chunks with offset", ra, "LAPIS-58"); print("    fs_read calls:", sum(n.endswith("fs_read") for n in names))
ra, names, errs = ask("Read scan.pdf and retell what it says.")
expect("scan: honestly says there is no text", ra.lower(), "text") if "scan" not in ra.lower() else expect("scan: honestly says it is a scan", ra.lower(), "scan")
print("\n##### SUMMARY:", "; ".join(f"{l}: {'OK' if ok else 'FAIL'}" for l, ok in verdicts))
print("ALL OK" if all(ok for _, ok in verdicts) else "SOME FAILED")
w.close()
