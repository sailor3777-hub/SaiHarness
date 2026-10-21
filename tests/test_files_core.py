import os, sys, pathlib, tempfile, shutil
from core import db, files, llm
from core.files import Selection, load_context, expand_selection, estimate_tokens, CHECKED, PARTIAL, UNCHECKED
from core.context_builder import build_messages, format_files
db.DB_PATH = pathlib.Path(sys.argv[1])

ok_all = True
def check(name, cond, extra=""):
    global ok_all
    ok_all &= bool(cond)
    print(("OK   " if cond else "FAIL ") + name + (f"  [{extra}]" if extra and not cond else ""))

tmp = tempfile.mkdtemp()
root = os.path.join(tmp, "proj")
def w(rel, data, mode="w", enc="utf-8"):
    p = os.path.join(root, rel); os.makedirs(os.path.dirname(p), exist_ok=True)
    if "b" in mode: open(p, mode).write(data)
    else: open(p, mode, encoding=enc).write(data)
    return p
w("a.txt", "hello, world\n"); w("sub/b.py", "print('hi')\n"); w("sub/deep/c.md", "# heading\n")
w(".git/config", "[core]"); w("node_modules/x.js", "x"); w("img.png", b"\x89PNG\x00\x00", "wb")
w("big.txt", "x" * 400_000); w("bin.dat", b"abc\x00def", "wb"); w("empty.txt", "")
w("cp1251.txt", "Русский текст в кодировке cp1251", enc="cp1251"); w("utf16.txt", "Юникод и 16 бит", enc="utf-16")   # Russian on purpose: legacy encodings
P = lambda *p: os.path.join(root, *p)

# --- Selection ---
s = Selection()
check("empty: nothing selected", s.state(root) == UNCHECKED and s.is_empty())
s.toggle(root, True)
check("a checked folder: descendants inherit", s.decision(P("sub", "b.py")) and s.state(root) == CHECKED and s.state(P("sub")) == CHECKED)
s.toggle(P("sub", "deep"), False)
check("an exclusion inside: the folder and ancestors are partial", s.state(P("sub", "deep")) == UNCHECKED and s.state(P("sub")) == PARTIAL and s.state(root) == PARTIAL)
check("an excluded node is out, its neighbors are in", not s.decision(P("sub", "deep", "c.md")) and s.decision(P("sub", "b.py")))
s.toggle(P("sub", "deep", "c.md"), True)
check("an explicit file inside an excluded folder is in", s.decision(P("sub", "deep", "c.md")) and s.state(P("sub", "deep")) == PARTIAL)
s.toggle(P("sub"), True)
check("checking a folder again resets exclusions inside", not s.exclude and s.state(P("sub")) == CHECKED)
s.toggle(root, False)
check("unchecking the root clears everything", s.is_empty() and not s.exclude and s.state(root) == UNCHECKED)
s.toggle(P("sub", "b.py"), True)
check("checking a file: ancestors are partial", s.state(P("sub")) == PARTIAL and s.state(root) == PARTIAL and s.state(P("a.txt")) == UNCHECKED)
s2 = Selection({root + "x", P("a.txt")}, {"C:\\other\\z"}); s2.prune([root])
check("prune removes paths outside the roots (and a similar prefix)", s2.include == {P("a.txt")} and not s2.exclude)

# --- loading ---
sel = Selection({root})
b = load_context(sel, [root])
names = [n for n, _ in b.files]
check("content: text and different encodings", names == ["proj/a.txt", "proj/cp1251.txt", "proj/empty.txt", "proj/img.png", "proj/utf16.txt", "proj/sub/b.py", "proj/sub/deep/c.md"], names)
txt = dict(b.files)
check("cp1251 read correctly", txt["proj/cp1251.txt"] == "Русский текст в кодировке cp1251")
check("utf-16 read correctly", txt["proj/utf16.txt"] == "Юникод и 16 бит")
check("service folders and binary extensions are not taken", not any(".git" in n or "node_modules" in n for n in names) and "vision_describe" in txt["proj/img.png"] and "img.png" in txt["proj/img.png"])
check("skipped: too large and binary (NUL)", sorted(r for _, r in b.skipped) == sorted(["too large (390 KB)", "binary"]) and len(b.skipped) == 2, b.skipped)
check("not truncated", not b.truncated and b.tokens > 0 and b.chars > 0)
sel.toggle(P("sub"), False)
check("an excluded folder is not read", [n for n, _ in load_context(sel, [root]).files if n.startswith("proj/sub")] == [])
sel.toggle(P("sub", "deep", "c.md"), True)
check("an explicit file from an excluded folder is read", "proj/sub/deep/c.md" in [n for n, _ in load_context(sel, [root]).files])
f_only = load_context(Selection({P("a.txt")}), [root])
check("one file", [n for n, _ in f_only.files] == ["proj/a.txt"])
old = files.MAX_FILES; files.MAX_FILES = 3
tr = load_context(Selection({root}), [root]); files.MAX_FILES = old
check("file count limit -> truncated", tr.truncated and len(tr.files) <= 3)
old = files.MAX_TOTAL_CHARS; files.MAX_TOTAL_CHARS = 30
tr = load_context(Selection({root}), [root]); files.MAX_TOTAL_CHARS = old
check("size limit -> truncated", tr.truncated and sum(len(t) for _, t in tr.files) <= 30)
check("a vanished file does not break loading", load_context(Selection({P("nope.txt")}), [root]).files == [])
check("estimate: Latin ≈ 5 chars/token (fitted to measurements)", abs(estimate_tokens("a" * 360) - 72) <= 1, estimate_tokens("a" * 360))
check("estimate: Cyrillic ≈ 3.2 chars/token (fitted to measurements)", abs(estimate_tokens("я" * 200) - 62) <= 1, estimate_tokens("я" * 200))
check("estimate: empty", estimate_tokens("") == 0)

# --- building the request ---
prof = {"system_prompt": "SYS", "memory": "MEM"}
m = build_messages(prof, [{"role": "user", "content": "q"}], b.files[:1])
sysmsg = m[0]["content"]
check("files in system after the prompt and memory", sysmsg.startswith("SYS\n\nMemory:\nMEM\n\nContext files") and "=== proj/a.txt ===\n```\nhello, world\n```" in sysmsg, sysmsg)
check("history after system", [x["role"] for x in m] == ["system", "user"])
check("without files there is no block", "Context files" not in build_messages(prof, [])[0]["content"])
check("a file with ``` is wrapped in ````", "````\na ```x``` b\n````" in format_files([("f", "a ```x``` b")]))
only_files = build_messages({"system_prompt": "", "memory": ""}, [], b.files[:1])
check("files only: system without empty sections", only_files[0]["content"].startswith("Context files"))

# --- DB ---
db.add_root(root); db.add_root(root)
check("roots: adding is idempotent", db.list_roots() == [root])
cid = db.create_chat()
db.save_chat_context(cid, {P("a.txt"), root}, {P("sub")})
inc, exc = db.load_chat_context(cid)
check("the chat's selection is saved and read", inc == {P("a.txt"), root} and exc == {P("sub")})
db.save_chat_context(cid, {P("a.txt")}, set())
check("saving replaces the previous selection", db.load_chat_context(cid) == ({P("a.txt")}, set()))
cid2 = db.create_chat(); db.save_chat_context(cid2, {root}, set())
check("chat selections are independent", db.load_chat_context(cid) == ({P("a.txt")}, set()) and db.load_chat_context(cid2) == ({root}, set()))
db.delete_chat(cid)
check("the selection is deleted with the chat (cascade)", db.load_chat_context(cid) == (set(), set()) and db.load_chat_context(cid2) == ({root}, set()))
db.remove_root(root)
check("the root is deleted", db.list_roots() == [])

# --- context limit of a real LM Studio ---
url = "http://127.0.0.1:1234/v1"
lim, loaded, mx = llm.get_context_info(url, "", "gemma-4-e4b-it-erotica")
check("loaded model: the working limit is known", loaded and isinstance(lim, int) and lim > 0 and mx and mx >= lim, (lim, loaded, mx))
print("      gemma-4-e4b-it-erotica: limit", lim, "maximum", mx)
import httpx
_models = httpx.get("http://127.0.0.1:1234/api/v0/models", timeout=5, trust_env=False).json()["data"]
_unloaded = next(m["id"] for m in _models if m.get("state") != "loaded" and m.get("type") != "embeddings")   # any unloaded chat model
lim2, loaded2, mx2 = llm.get_context_info(url, "", _unloaded)
check("a model that is not loaded: limit unknown, maximum known", lim2 is None and not loaded2 and mx2, (lim2, loaded2, mx2))
check("unknown model", llm.get_context_info(url, "", "no-such") == (None, False, None))
check("server unavailable -> a quiet refusal", llm.get_context_info("http://127.0.0.1:1/v1", "", "x") == (None, False, None))

# --- documents (PDF / Word / Excel) go into the context as text ---
import docx as _docx, openpyxl as _xl
from pypdf import PdfWriter as _PdfWriter
from core import docs as _docs
droot = os.path.join(tmp, "docs"); os.makedirs(droot)
D = lambda n: os.path.join(droot, n)
_d = _docx.Document(); _d.add_paragraph("Lease agreement #77"); _d.save(D("contract.docx"))
_wb = _xl.Workbook(); _wb.active.append(["Item", 100]); _wb.save(D("book.xlsx"))
_w = _PdfWriter(); _w.add_blank_page(100, 100)
with open(D("scan.pdf"), "wb") as _f: _w.write(_f)
dctx = load_context(Selection({D("contract.docx"), D("book.xlsx"), D("scan.pdf")}), [droot])
dt = dict(dctx.files)
check("docx in the context as text", "Lease agreement #77" in dt["docs/contract.docx"] and "binary" not in dt["docs/contract.docx"].lower())
check("xlsx in the context as text", "Item\t100" in dt["docs/book.xlsx"], dt.get("docs/book.xlsx"))
check("a scan is skipped with a reason, not a placeholder", "docs/scan.pdf" not in dt and any("no text layer" in r for _, r in dctx.skipped), dctx.skipped)
check("document tokens are counted", dctx.tokens > 0)
folder = load_context(Selection({droot}), [droot])
check("a checked folder does not pull in documents by itself (only an explicit check)", folder.files == [] and folder.skipped == [])
_calls = []
_orig = _docs.read_document
_docs.read_document = lambda path: (_calls.append(path), _orig(path))[1]
_docs._cache.clear()
load_context(Selection({D("contract.docx")}), [droot]); load_context(Selection({D("contract.docx")}), [droot])
check("reading a document again comes from the cache", len(_calls) == 1, _calls)
_d2 = _docx.Document(); _d2.add_paragraph("New version"); _d2.save(D("contract.docx")); os.utime(D("contract.docx"), (1, 2_000_000_000))
check("a changed file is read again", "New version" in dict(load_context(Selection({D("contract.docx")}), [droot]).files)["docs/contract.docx"] and len(_calls) == 2, _calls)
_docs.read_document = _orig
_old = files.MAX_FILE_BYTES; files.MAX_FILE_BYTES = 5
_big = load_context(Selection({D("book.xlsx")}), [droot]); files.MAX_FILE_BYTES = _old
check("too long a document text is skipped", _big.files == [] and "document text too large" in _big.skipped[0][1], _big.skipped)

# a text file with rare NULs (like the Windows DumpStack.log: NUL in the header and zero padding up to 12 KB)
_dump = b"DLOGFILE00010000DUMPC%" + b"\x00" * 10 + b"\nBugCheckCode 0x19C\n" * 100 + b"\x00" * 3000
_t, _r = files.read_text_file(w("DumpStack.log", _dump, "wb"))
check("text with rare NULs is read, the NULs are cut out", _r == "" and _t and "BugCheckCode 0x19C" in _t and "\x00" not in _t and _t.startswith("DLOGFILE00010000DUMPC%"), (_t or "")[:40] + _r)
_t, _r = files.read_text_file(w("bin2.dat", b"MZ\x90\x00\x03\x00\x00\x00\x04\x00\x00\x00" * 50, "wb"))
check("a real binary is still skipped", _t is None and _r == "binary", _r)

shutil.rmtree(tmp, ignore_errors=True)
print("ALL OK" if ok_all else "SOME FAILED")

